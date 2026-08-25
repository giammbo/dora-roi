"""FIELD_CATALOG against the official EBA files.

Golden rule 6 says the EBA annotated template and sample package win over this
repository's reconstructed tables. A one-off reconciliation would satisfy that
once and then rot, so it lives here instead: the official files are vendored in
`tests/fixtures/eba/` and the suite re-derives the comparison on every run.

If the EBA publishes a revision and someone drops it in, these tests are what
tells them what changed.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

from dora_roi.models.enums import (
    ContractualArrangementType,
    DataSensitiveness,
    ICTServiceType,
    IdentifierType,
    ImpactOfDiscontinuing,
    LevelOfReliance,
    ReasonNotSubstitutable,
    Reintegration,
    Substitutability,
    TypeOfPerson,
)
from dora_roi.models.templates import FIELD_CATALOG, UNVERIFIED_FIELD_NAMES

EBA = Path(__file__).parent / "fixtures" / "eba"
ANNOTATED = EBA / "annotated_dora_4.0.xlsx"
SAMPLE = EBA / "sample_dora_package.zip"

_CODE = re.compile(r"^\d{4}$")


def official_layout() -> dict[str, dict[str, str]]:
    """code -> name per template, straight out of the EBA annotated workbook."""
    workbook = openpyxl.load_workbook(ANNOTATED, read_only=True, data_only=True)
    layout: dict[str, dict[str, str]] = {}
    for sheet in workbook.sheetnames:
        if not sheet.startswith("B_"):
            continue
        rows = list(workbook[sheet].iter_rows(max_row=12, max_col=40, values_only=True))
        for index, row in enumerate(rows):
            cells = [str(cell).strip() if cell is not None else "" for cell in row]
            codes = [(column, cell) for column, cell in enumerate(cells) if _CODE.match(cell)]
            if len(codes) >= 2:
                names = rows[index - 1]
                layout[sheet] = {code: re.sub(r"\s+", " ", str(names[column]).strip()) for column, code in codes}
                break
    return layout


def sample_headers() -> dict[str, list[str]]:
    """template -> column codes, from the official sample report package."""
    headers: dict[str, list[str]] = {}
    with zipfile.ZipFile(SAMPLE) as package:
        for name in package.namelist():
            stem = Path(name).name
            if not (stem.startswith("b_") and stem.endswith(".csv")):
                continue
            text = package.read(name).decode("utf-8-sig")
            row = next(csv.reader(io.StringIO(text)))
            headers[stem[:-4].upper()] = [column[1:] for column in row if column.startswith("c")]
    return headers


@pytest.fixture(scope="module")
def layout() -> dict[str, dict[str, str]]:
    return official_layout()


@pytest.fixture(scope="module")
def sample() -> dict[str, list[str]]:
    return sample_headers()


class TestTheFixturesAreWhatTheyClaim:
    def test_the_annotated_workbook_has_a_sheet_per_template(self, layout) -> None:
        assert set(layout) == set(FIELD_CATALOG)

    def test_the_sample_package_has_a_csv_per_template(self, sample) -> None:
        assert set(sample) == set(FIELD_CATALOG)

    def test_the_sample_package_carries_the_filing_scaffolding(self) -> None:
        with zipfile.ZipFile(SAMPLE) as package:
            names = {Path(n).name for n in package.namelist()}
        assert {"report.json", "parameters.csv", "FilingIndicators.csv", "reportPackage.json"} <= names


class TestCatalogMatchesTheOfficialLayout:
    @pytest.mark.parametrize("template", sorted(FIELD_CATALOG))
    def test_codes_match(self, template: str, layout) -> None:
        assert {s.code for s in FIELD_CATALOG[template]} == set(layout[template])

    @pytest.mark.parametrize("template", sorted(FIELD_CATALOG))
    def test_names_match_verbatim(self, template: str, layout) -> None:
        assert {s.code: s.name for s in FIELD_CATALOG[template]} == layout[template]

    def test_b_99_01_is_one_row_of_nineteen_definition_columns(self, layout) -> None:
        """The template the drafting notes got wrong, now matching the official layout."""
        assert len(layout["B_99.01"]) == 19
        assert {s.code for s in FIELD_CATALOG["B_99.01"]} == set(layout["B_99.01"])
        assert layout["B_99.01"]["0010"] == "Standalone arrangement"

    def test_every_template_is_now_reconciled(self) -> None:
        assert UNVERIFIED_FIELD_NAMES == frozenset()


class TestCatalogMatchesTheSamplePackage:
    """The annotated layout and the sample package are independent sources."""

    @pytest.mark.parametrize("template", sorted(FIELD_CATALOG))
    def test_codes_match(self, template: str, sample) -> None:
        assert {s.code for s in FIELD_CATALOG[template]} == set(sample[template])

    def test_the_two_official_sources_agree_with_each_other(self, layout, sample) -> None:
        for template in FIELD_CATALOG:
            assert set(layout[template]) == set(sample[template]), template


class TestFilingFormatFacts:
    """Facts about the wire format the exporter has to honour, captured now."""

    def test_columns_are_prefixed_with_c(self) -> None:
        with zipfile.ZipFile(SAMPLE) as package:
            name = next(n for n in package.namelist() if n.endswith("b_05.01.csv"))
            header = next(csv.reader(io.StringIO(package.read(name).decode("utf-8-sig"))))
        assert all(column.startswith("c") for column in header)
        assert "0010" not in header and "c0010" in header

    def test_csv_files_are_lower_case_and_indicators_are_upper(self) -> None:
        with zipfile.ZipFile(SAMPLE) as package:
            names = {Path(n).name for n in package.namelist() if n.endswith(".csv")}
            indicators = package.read(next(n for n in package.namelist() if n.endswith("FilingIndicators.csv"))).decode(
                "utf-8-sig"
            )
        assert "b_05.01.csv" in names
        assert "B_05.01,true" in indicators

    def test_closed_lists_are_eba_prefixed_qnames_not_labels(self) -> None:
        """The enums hold human labels; a filing needs `eba_XX:code`."""
        with zipfile.ZipFile(SAMPLE) as package:
            name = next(n for n in package.namelist() if n.endswith("b_05.01.csv"))
            rows = list(csv.DictReader(io.StringIO(package.read(name).decode("utf-8-sig"))))
        values = " ".join(rows[0].values())
        assert "eba_" in values and ":" in values

    def test_report_json_extends_the_dora_4_0_module(self) -> None:
        with zipfile.ZipFile(SAMPLE) as package:
            name = next(n for n in package.namelist() if n.endswith("reports/report.json"))
            report = json.loads(package.read(name))
        assert report["documentInfo"]["extends"] == [
            "http://www.eba.europa.eu/eu/fr/xbrl/crr/fws/dora/4.0/mod/dora.json"
        ]

    def test_the_package_wraps_everything_in_one_directory(self) -> None:
        """A detail the drafting notes never mentioned, and one a validator rejects on.

        META-INF/ and reports/ do not sit at the root of the zip: they sit
        inside a single directory named exactly like the package.
        """
        with zipfile.ZipFile(SAMPLE) as package:
            roots = {Path(n).parts[0] for n in package.namelist()}
            inner = {Path(n).parts[1] for n in package.namelist() if len(Path(n).parts) > 1}
        assert len(roots) == 1
        assert inner == {"META-INF", "reports"}

    def test_the_zip_name_follows_the_documented_convention(self) -> None:
        with zipfile.ZipFile(SAMPLE) as package:
            root = Path(package.namelist()[0]).parts[0]
        assert re.match(r"^[A-Z0-9]{20}\.(CON|IND)_[A-Z]{2}_DORA010100_DORA_\d{4}-\d{2}-\d{2}_\d+$", root)


@pytest.fixture(scope="module")
def official() -> dict[str, dict[str, str]]:
    """label -> qname per list, read from the workbook, not from our own file."""
    workbook = openpyxl.load_workbook(EBA / "possible_values.xlsx", read_only=True, data_only=True)
    by_list: dict[str, dict[str, str]] = {}
    for sheet in workbook.sheetnames:
        rows = list(workbook[sheet].iter_rows(values_only=True))
        for index, row in enumerate(rows[:12]):
            cells = [str(cell).strip() if cell is not None else "" for cell in row]
            heads = [(column, cell) for column, cell in enumerate(cells) if cell.startswith("LIST")]
            if not heads:
                continue
            for column, name in heads:
                values = by_list.setdefault(f"{sheet}.{name}", {})
                for data in rows[index + 1 :]:
                    if len(data) > column + 1 and data[column] not in (None, ""):
                        values.setdefault(str(data[column + 1]).strip(), str(data[column]).strip())
            break
    return by_list


class TestClosedListsMatchTheOfficialValues:
    """Enum labels against the EBA list of possible values."""

    ENUMS = [
        (ICTServiceType, "ict_service_type"),
        (IdentifierType, "identifier_type"),
        (TypeOfPerson, "type_of_person"),
        (ContractualArrangementType, "contractual_arrangement"),
        (DataSensitiveness, "data_sensitiveness"),
        (LevelOfReliance, "level_of_reliance"),
        (Substitutability, "substitutability"),
        (ReasonNotSubstitutable, "reason_not_substitutable"),
        (Reintegration, "reintegration"),
        (ImpactOfDiscontinuing, "impact_of_discontinuing"),
    ]

    @pytest.mark.parametrize(("enum", "list_name"), ENUMS, ids=[e.__name__ for e, _ in ENUMS])
    def test_every_member_resolves_to_a_qname(self, enum, list_name: str) -> None:
        """A label that does not resolve is a label we transcribed wrong."""
        for member in enum:
            assert member.qname.startswith("eba_"), f"{enum.__name__}.{member.name}"

    @pytest.mark.parametrize(("enum", "list_name"), ENUMS, ids=[e.__name__ for e, _ in ENUMS])
    def test_the_packaged_file_matches_the_workbook(self, enum, list_name, official) -> None:
        packaged = json.loads((Path("src/dora_roi/data") / "eba_closed_lists.json").read_text(encoding="utf-8"))[
            list_name
        ]
        from_workbook = official[packaged["source"]]
        assert {v["label"]: v["qname"] for v in packaged["values"]} == from_workbook

    def test_qnames_are_distinct_within_a_list(self) -> None:
        for enum, _ in self.ENUMS:
            qnames = [member.qname for member in enum]
            assert len(set(qnames)) == len(qnames), enum.__name__

    def test_the_service_qname_keeps_its_s_code(self) -> None:
        for service in ICTServiceType:
            assert service.qname == f"eba_TA:{service.value}"

    def test_a_short_key_still_carries_the_official_wording(self) -> None:
        assert IdentifierType.LEI.label == "Legal Entity Identfier (LEI)"
        assert IdentifierType.NIN.label == "National code"

    def test_the_lists_the_notes_never_enumerated_are_now_packaged(self) -> None:
        packaged = json.loads((Path("src/dora_roi/data") / "eba_closed_lists.json").read_text(encoding="utf-8"))
        assert len(packaged["entity_type"]["values"]) == 22  # the notes said "22 types"
        assert len(packaged["entity_type_b0102"]["values"]) == 24  # "+2 extra values in B_01.02"
        assert packaged["hierarchy"]["values"]
        assert packaged["licenced_activity"]["values"]
